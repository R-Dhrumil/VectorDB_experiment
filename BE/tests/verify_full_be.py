import io
import time
import unittest
import openpyxl
from docx import Document as DocxDocument
from pypdf import PdfWriter
from fastapi.testclient import TestClient

from main import app
from app.extractors.pdf import extract_pdf
from app.extractors.docx import extract_docx
from app.extractors.excel import extract_excel
from app.extractors.txt import extract_txt
from app.extractors.factory import extract_document_blocks
from app.processing.cleaner import clean_text
from app.processing.chunker import chunk_document_blocks
from app.database.vector_store import get_vector_store, list_documents, delete_document_by_id

client = TestClient(app)

class TestFullBackendVerification(unittest.TestCase):

    def test_01_cleaner_and_extractors(self):
        print("\n[1/5] Testing Text Cleaner and Extractors...")
        # 1. Cleaner
        raw_dirty = "  Header   \n\n\n\nBody text   with   spaces.\x00  "
        cleaned = clean_text(raw_dirty)
        self.assertEqual(cleaned, "Header\n\nBody text with spaces.")

        # 2. TXT Extractor
        txt_bytes = b"Hello World!\nThis is a sample text file for RAG testing."
        txt_blocks = extract_txt(txt_bytes)
        self.assertEqual(len(txt_blocks), 1)
        self.assertIn("Hello World!", txt_blocks[0]["text"])

        # 3. DOCX Extractor
        docx_io = io.BytesIO()
        doc = DocxDocument()
        doc.add_heading("Section 1: RAG System", level=1)
        doc.add_paragraph("Vector database enables semantic search across enterprise documents.")
        doc.save(docx_io)
        docx_bytes = docx_io.getvalue()

        docx_blocks = extract_docx(docx_bytes)
        self.assertGreater(len(docx_blocks), 0)
        self.assertIn("Vector database", docx_blocks[1]["text"])

        # 4. XLSX Extractor
        xlsx_io = io.BytesIO()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Products"
        ws.append(["Item", "Category", "Price"])
        ws.append(["Laptop", "Electronics", 1200])
        ws.append(["Chair", "Furniture", 150])
        wb.save(xlsx_io)
        xlsx_bytes = xlsx_io.getvalue()

        xlsx_blocks = extract_excel(xlsx_bytes)
        self.assertEqual(len(xlsx_blocks), 1)
        self.assertIn("Item: Laptop", xlsx_blocks[0]["text"])
        self.assertIn("Category: Electronics", xlsx_blocks[0]["text"])

        # 5. Factory Dispatcher
        factory_pdf = extract_document_blocks(txt_bytes, "doc.txt")
        self.assertEqual(len(factory_pdf), 1)
        print("  -> All extractors (TXT, DOCX, XLSX) passed!")

    def test_02_chunker_strategies(self):
        print("[2/5] Testing Chunking Strategies...")
        sample_blocks = [{"text": "Paragraph 1 sentence 1. Sentence 2.\n\nParagraph 2 sentence 3. Sentence 4.", "metadata": {}}]

        for strat in ["recursive", "fixed", "sentence", "paragraph"]:
            chunks = chunk_document_blocks(
                blocks=sample_blocks,
                doc_id="test_doc",
                file_name="sample.txt",
                strategy=strat,
                chunk_size=50,
                chunk_overlap=10
            )
            self.assertGreater(len(chunks), 0)
            self.assertEqual(chunks[0].metadata["file_name"], "sample.txt")
        print("  -> Chunking strategies (recursive, fixed, sentence, paragraph) passed!")

    def test_03_vector_store_operations(self):
        print("[3/5] Testing ChromaDB Vector Store & Embeddings...")
        vs = get_vector_store()
        self.assertIsNotNone(vs)

        sample_blocks = [{"text": "ChromaDB vector database stores high-dimensional dense embeddings.", "metadata": {}}]
        chunks = chunk_document_blocks(sample_blocks, doc_id="vector_test_doc", file_name="vector.txt")
        vs.add_documents(chunks)

        search_results = vs.similarity_search("What does ChromaDB store?", k=1)
        self.assertGreater(len(search_results), 0)
        self.assertIn("ChromaDB", search_results[0].page_content)
        print("  -> Vector Store embedding & search passed!")

    def test_04_fastapi_endpoints(self):
        print("[4/5] Testing FastAPI REST Endpoints...")
        # 1. Upload DOCX File
        docx_io = io.BytesIO()
        doc = DocxDocument()
        doc.add_heading("FastAPI Verification Document", level=1)
        doc.add_paragraph("FastAPI provides automatic OpenAPI interactive documentation at /docs.")
        doc.save(docx_io)

        response = client.post(
            "/documents/upload",
            files={"file": ("fastapi_doc.docx", docx_io.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            data={"strategy": "recursive", "chunk_size": 300, "chunk_overlap": 30}
        )
        self.assertEqual(response.status_code, 200)
        upload_res = response.json()
        task_id = upload_res["task_id"]
        doc_id = upload_res["doc_id"]

        # 2. Poll Task Progress
        max_attempts = 10
        completed = False
        for _ in range(max_attempts):
            time.sleep(0.5)
            status_res = client.get(f"/documents/tasks/{task_id}").json()
            if status_res.get("status") == "COMPLETED":
                completed = True
                break

        self.assertTrue(completed, f"Background task did not complete: {status_res}")

        # 3. Query Endpoint
        query_res = client.post(
            "/documents/query",
            json={"query": "Where is the documentation located?", "top_k": 2}
        ).json()
        self.assertGreater(query_res["results_count"], 0)
        self.assertIn("/docs", query_res["results"][0]["content"])

        # 4. List Documents Endpoint
        docs_res = client.get("/documents").json()
        doc_ids = [d["doc_id"] for d in docs_res["documents"]]
        self.assertIn(doc_id, doc_ids)

        # 5. Delete Document Endpoint
        del_res = client.delete(f"/documents/{doc_id}").json()
        self.assertIn("Successfully deleted", del_res["message"])
        print("  -> FastAPI endpoints (/upload, /tasks, /query, /documents, DELETE) passed!")

    def test_05_legacy_endpoints(self):
        print("[5/5] Testing Legacy Endpoint Backwards Compatibility...")
        store_res = client.post("/store", json={"text": "Legacy endpoint test content.", "strategy": "recursive"}).json()
        self.assertIn("Successfully embedded", store_res["message"])

        query_res = client.post("/query", json={"query": "Legacy endpoint", "top_k": 1}).json()
        self.assertGreater(len(query_res["results"]), 0)
        print("  -> Legacy endpoints (/store, /query) passed!")


if __name__ == "__main__":
    unittest.main()
