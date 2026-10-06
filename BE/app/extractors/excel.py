import io
import openpyxl

def extract_excel(file_bytes: bytes) -> list[dict]:
    """
    Extracts structured key-value data from XLSX files sheet by sheet.
    Preserves column headers so context is maintained per row chunk.
    """
    excel_file = io.BytesIO(file_bytes)
    wb = openpyxl.load_workbook(excel_file, data_only=True)
    blocks = []

    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue

        headers = [str(h or f"Column_{i+1}") for i, h in enumerate(rows[0])]
        sheet_text_lines = []

        for row_idx, row in enumerate(rows[1:], start=2):
            row_pairs = []
            for col_idx, cell_value in enumerate(row):
                if cell_value is not None:
                    header = headers[col_idx] if col_idx < len(headers) else f"Column_{col_idx+1}"
                    row_pairs.append(f"{header}: {cell_value}")

            if row_pairs:
                row_str = f"Sheet '{sheet_name}' (Row {row_idx}): " + " | ".join(row_pairs)
                sheet_text_lines.append(row_str)

        if sheet_text_lines:
            blocks.append({
                "text": "\n".join(sheet_text_lines),
                "metadata": {
                    "sheet_name": sheet_name,
                    "row_count": len(rows) - 1
                }
            })

    return blocks
